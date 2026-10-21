import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import create_app
from app.models import User, db
from tasks.persistence.vms import save_networks, save_vms
from werkzeug.security import generate_password_hash


@pytest.fixture
def app_client(monkeypatch):
    monkeypatch.setenv("APP_DATABASE_URL", "sqlite:///:memory:")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///:memory:")
    app = create_app()
    app.config.update(TESTING=True)
    with app.app_context():
        save_vms(
            [{"VMId": "vm-1", "Name": "app-01", "State": "Running"}],
            "HOST-01",
            cluster_name="CLUSTER-01",
        )
        save_networks(
            [
                {
                    "VMId": "vm-1",
                    "AdapterName": "Production",
                    "SwitchName": "vSwitch-Prod",
                    "MacAddress": "00155D000001",
                    "IPAddresses": "10.0.0.5",
                    "IsConnected": True,
                },
                {
                    "VMId": "vm-1",
                    "AdapterName": "Backup",
                    "SwitchName": "vSwitch-Backup",
                    "MacAddress": "00155D000002",
                    "IPAddresses": "10.0.0.5,10.0.0.6",
                    "IsConnected": True,
                },
                {
                    "VMId": "vm-1",
                    "AdapterName": "Unconfigured",
                    "SwitchName": "vSwitch-Prod",
                    "MacAddress": "00155D000003",
                    "IPAddresses": "",
                    "IsConnected": False,
                },
            ],
            cluster_name="CLUSTER-01",
        )
        db.session.add(
            User(
                username="export-tester",
                password_hash=generate_password_hash("export-test-password"),
            )
        )
        db.session.commit()

        client = app.test_client()
        response = client.post(
            "/login",
            data={"username": "export-tester", "password": "export-test-password"},
        )
        assert response.status_code == 302
        yield client


def _workbook_rows(response):
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook["VMs"]
    rows = list(sheet.iter_rows(values_only=True))
    return rows[0], rows[1:]


def test_export_vms_adds_aggregated_mac_addresses(app_client):
    headers, rows = _workbook_rows(app_client.get("/export/vms.xlsx"))

    assert "IP Addresses" in headers
    assert "MAC Addresses" in headers
    assert headers.index("MAC Addresses") == headers.index("IP Addresses") + 1

    row = dict(zip(headers, rows[0]))
    assert set(str(row["IP Addresses"]).split(",")) == {"10.0.0.5", "10.0.0.6"}
    assert set(str(row["MAC Addresses"]).split(",")) == {
        "00155D000001",
        "00155D000002",
        "00155D000003",
    }


def test_filtered_export_vms_adds_aggregated_mac_addresses(app_client):
    response = app_client.post(
        "/export/vms/filtered.xlsx",
        data={"vm_ids": json.dumps(["vm-1"])},
    )
    headers, rows = _workbook_rows(response)

    assert "IP Addresses" in headers
    assert "MAC Addresses" in headers
    assert headers.index("MAC Addresses") == headers.index("IP Addresses") + 1

    row = dict(zip(headers, rows[0]))
    assert set(str(row["IP Addresses"]).split(",")) == {"10.0.0.5", "10.0.0.6"}
    assert set(str(row["MAC Addresses"]).split(",")) == {
        "00155D000001",
        "00155D000002",
        "00155D000003",
    }
